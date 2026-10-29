"""Contas lidas da planilha do financeiro (Financeiro Arapuca 2K26 no Google Drive).

O financeiro lança tudo na planilha; o app baixa o arquivo, lê e mostra. Nada é lançado no app.

O que o app lê (procurando pelos títulos das colunas, não pela posição):
- uma aba por mês (Janeiro, Fevereiro...): a tabela com "Morador | Casa | Despesas | Outros |
  Pendências | Caixinha | Desconto | Aluguel | Pagou | Saldo" e os quadros de gastos que têm
  "Data | Descrição | ... | Valor" (Mensalidades, Cartão de Crédito, Outros Gastos);
- a aba Créditos: "Morador | Saldo" (negativo = deve à rep).

Acesso ao arquivo, nesta ordem:
1. conta de serviço do Google (credenciais/google.json): a planilha fica privada, compartilhada só
   com o e-mail da conta de serviço;
2. sem credenciais: o link público de download (a planilha precisa estar como "qualquer pessoa com o link").
"""
import io
import json
import os
import re
import threading
import urllib.parse
import urllib.request
from datetime import datetime, timedelta
from pathlib import Path

from modulos.planilha import normalizar

MESES = ["Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho", "Julho", "Agosto",
         "Setembro", "Outubro", "Novembro", "Dezembro"]
COLUNAS_MORADOR = {"casa": "casa", "despesas": "despesas", "outros": "outros", "pendencias": "pendencias",
                   "caixinha": "caixinha", "desconto": "desconto", "aluguel": "total", "pagou": "pagou",
                   "saldo": "saldo"}
INTERVALO = timedelta(hours=1)
CREDENCIAIS_PADRAO = Path("credenciais") / "google.json"
_trava = threading.Lock()


# ---------- leitura do arquivo ----------

def texto(valor):
    return str(valor).strip() if valor is not None else ""


def numero(valor):
    if isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return float(valor)
    return None


def centavos(valor):
    n = numero(valor)
    return round(n * 100) if n is not None else 0


def ler_ano(wb):
    """As abas têm o título "FINANCEIRO 2K26" (o Resumo às vezes ficou com um ano velho, então vale o maior)."""
    anos = []
    for ws in wb.worksheets:
        for linha in ws.iter_rows(max_row=3, max_col=10, values_only=True):
            for valor in linha:
                achou = re.search(r"2K(\d{2})", texto(valor).upper())
                if achou:
                    anos.append(2000 + int(achou.group(1)))
    return max(anos) if anos else datetime.now().year


def ler_moradores(linhas):
    """A tabela "Morador | Casa | ... | Aluguel | Pagou | Saldo" da aba do mês."""
    for i, linha in enumerate(linhas):
        nomes = [normalizar(texto(v)) for v in linha]
        if "morador" not in nomes or "aluguel" not in nomes:
            continue
        col_morador = nomes.index("morador")
        col_quarto = nomes.index("quarto") if "quarto" in nomes else None
        colunas = {}
        for j, n in enumerate(nomes):
            if j > col_morador and n in COLUNAS_MORADOR and COLUNAS_MORADOR[n] not in colunas:
                colunas[COLUNAS_MORADOR[n]] = j  # o primeiro: mais à direita tem outro quadro "Despesas"
        resultado, quarto = [], ""
        for abaixo in linhas[i + 1:]:
            nome = texto(abaixo[col_morador]) if col_morador < len(abaixo) else ""
            if col_quarto is not None and col_quarto < len(abaixo) and texto(abaixo[col_quarto]):
                quarto = texto(abaixo[col_quarto])
            if normalizar(nome) in ("total", "---"):
                break
            if not nome:
                if not any(texto(v) for v in abaixo[:col_morador + 10]):
                    break
                continue
            valores = {campo: centavos(abaixo[j]) if j < len(abaixo) else 0 for campo, j in colunas.items()}
            resultado.append({"nome": nome, "quarto": quarto, **valores})
        return resultado
    return []


def ler_quadros(linhas):
    """Quadros de gastos: cabeçalho com "Descrição" ... "Valor"; o título fica logo acima."""
    quadros = []
    for i, linha in enumerate(linhas):
        nomes = [normalizar(texto(v)) for v in linha]
        for j, nome in enumerate(nomes):
            if nome != "descricao":
                continue
            inicio = j - 1 if j > 0 and nomes[j - 1] == "data" else j
            try:
                fim = next(k for k in range(j + 1, min(j + 8, len(nomes))) if nomes[k] == "valor")
            except StopIteration:
                continue
            campos = {nomes[k]: k for k in range(inicio, fim + 1) if nomes[k]}
            titulo = ""
            for acima in range(i - 1, max(i - 4, -1), -1):
                t = texto(linhas[acima][inicio]) if inicio < len(linhas[acima]) else ""
                if t and not normalizar(t).startswith("total"):
                    titulo = t
                    break
            itens, vazias = [], 0
            for abaixo in linhas[i + 1:]:
                celulas = [abaixo[k] if k < len(abaixo) else None for k in range(inicio, fim + 1)]
                if any(normalizar(texto(c)).startswith("total") for c in celulas):
                    break
                descricao = texto(abaixo[j]) if j < len(abaixo) else ""
                valor = numero(abaixo[fim]) if fim < len(abaixo) else None
                if not descricao and valor is None:
                    vazias += 1
                    if vazias >= 4:
                        break
                    continue
                vazias = 0
                if valor is None or not descricao:
                    continue
                pega = lambda chave: texto(abaixo[campos[chave]]) if chave in campos and campos[chave] < len(abaixo) else ""  # noqa: E731
                itens.append({"descricao": descricao, "valor": round(valor * 100), "data": pega("data"),
                              "status": pega("status"), "parcela": pega("parcela"),
                              "responsavel": pega("responsavel")})
            quadros.append({"titulo": titulo or "Gastos", "itens": itens,
                            "total": sum(it["valor"] for it in itens)})
    return quadros


def ler_creditos(linhas):
    for i, linha in enumerate(linhas):
        nomes = [normalizar(texto(v)) for v in linha]
        if "morador" in nomes and "saldo" in nomes:
            cm, cs = nomes.index("morador"), nomes.index("saldo")
            saldos = {}
            for abaixo in linhas[i + 1:]:
                nome = texto(abaixo[cm]) if cm < len(abaixo) else ""
                if not nome or normalizar(nome) == "total":
                    if saldos:
                        break
                    continue
                saldos[nome] = centavos(abaixo[cs] if cs < len(abaixo) else None)
            return saldos
    return {}


def ler_arquivo(conteudo, ano=None):
    """xlsx (bytes) -> {"ano", "meses": {"2026-10": {...}}, "creditos": {nome: centavos}}.
    ano: o ano da planilha (um arquivo por ano); sem ele, sai do título "FINANCEIRO 2K26" das abas."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(conteudo), data_only=True, read_only=True)
    ano = ano or ler_ano(wb)
    abas = {normalizar(ws.title): ws for ws in wb.worksheets}
    dados = {"ano": ano, "meses": {}, "creditos": {}}
    for numero_mes, nome in enumerate(MESES, start=1):
        ws = abas.get(normalizar(nome))
        if ws is None:
            continue
        linhas = [list(l) for l in ws.iter_rows(max_row=300, max_col=60, values_only=True)]
        moradores = ler_moradores(linhas)
        if moradores:
            dados["meses"][f"{ano}-{numero_mes:02d}"] = {"nome": nome, "moradores": moradores,
                                                        "quadros": ler_quadros(linhas)}
    if "creditos" in abas:
        dados["creditos"] = ler_creditos([list(l) for l in abas["creditos"].iter_rows(max_row=200, max_col=10,
                                                                                     values_only=True)])
    wb.close()
    return dados


# ---------- download ----------

def id_do_link(link):
    """Aceita o link do Drive/Planilhas ou só o id."""
    link = link.strip()
    achou = re.search(r"/d/([A-Za-z0-9_-]{20,})", link) or re.search(r"[?&]id=([A-Za-z0-9_-]{20,})", link)
    if achou:
        return achou.group(1)
    return link if re.fullmatch(r"[A-Za-z0-9_-]{20,}", link) else ""


def caminho_credenciais():
    return Path(os.environ.get("ARAPUCA_GOOGLE_CREDENCIAIS", CREDENCIAIS_PADRAO))


def token_conta_servico():
    from google.auth.transport.requests import Request
    from google.oauth2 import service_account

    cred = service_account.Credentials.from_service_account_file(
        str(caminho_credenciais()), scopes=["https://www.googleapis.com/auth/drive.readonly"])
    cred.refresh(Request())
    return cred.token


def email_conta_servico():
    try:
        return json.loads(caminho_credenciais().read_text()).get("client_email", "")
    except (OSError, ValueError):
        return ""


def _get(url, token=None, timeout=30):
    pedido = urllib.request.Request(url, headers={"Authorization": f"Bearer {token}"} if token else {})
    with urllib.request.urlopen(pedido, timeout=timeout) as resposta:
        return resposta.read()


def baixar(arquivo_id):
    """Devolve os bytes do xlsx. Planilha Google é exportada como xlsx; arquivo .xlsx vem como está."""
    api = "https://www.googleapis.com/drive/v3/files/" + urllib.parse.quote(arquivo_id)
    if caminho_credenciais().exists():
        token = token_conta_servico()
        meta = json.loads(_get(api + "?fields=mimeType&supportsAllDrives=true", token))
        if meta.get("mimeType") == "application/vnd.google-apps.spreadsheet":
            tipo = urllib.parse.quote("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            return _get(f"{api}/export?mimeType={tipo}", token)
        return _get(api + "?alt=media&supportsAllDrives=true", token)
    # sem credenciais: link público (planilha Google ou arquivo .xlsx compartilhado)
    try:
        return _get(f"https://docs.google.com/spreadsheets/d/{arquivo_id}/export?format=xlsx")
    except Exception:
        return _get(f"https://drive.google.com/uc?export=download&id={arquivo_id}")


# ---------- guardar no banco e atualizar ----------

def ler_config(conn, chave, padrao=""):
    linha = conn.execute("SELECT valor FROM config WHERE chave = ?", (chave,)).fetchone()
    return linha["valor"] if linha else padrao


def salvar_config(conn, **valores):
    conn.executemany("INSERT OR REPLACE INTO config (chave, valor) VALUES (?, ?)",
                     [(k, str(v)) for k, v in valores.items()])


def arquivo_configurado(conn):
    return ler_config(conn, "financeiro_arquivo_id")


def carregar(conn):
    """O que foi lido da planilha da última vez (fica salvo no banco, então abre sem internet)."""
    try:
        return json.loads(ler_config(conn, "financeiro_dados", "{}")) or {}
    except ValueError:
        return {}


def status(conn):
    em = ler_config(conn, "financeiro_sync_em")
    return {"em": datetime.fromisoformat(em) if em else None, "erro": ler_config(conn, "financeiro_sync_erro"),
            "arquivo": arquivo_configurado(conn), "email": email_conta_servico()}


def sincronizar(baixar_arquivo=None):
    baixar_arquivo = baixar_arquivo or baixar  # lido na hora (os testes trocam o baixar)
    from database import conectar_db

    with _trava:
        conn = conectar_db()
        try:
            arquivo_id = arquivo_configurado(conn)
            if not arquivo_id:
                return "Nenhuma planilha configurada"
            try:
                ano = ler_config(conn, "financeiro_ano")
                dados = ler_arquivo(baixar_arquivo(arquivo_id), int(ano) if ano.isdigit() else None)
                if not dados["meses"]:
                    raise ValueError("não achei nenhuma aba de mês com a tabela de moradores")
            except Exception as erro:
                salvar_config(conn, financeiro_sync_em=datetime.now().isoformat(timespec="minutes"),
                              financeiro_sync_erro=f"{type(erro).__name__}: {erro}"[:300])
                conn.commit()
                return str(erro)
            salvar_config(conn, financeiro_dados=json.dumps(dados, ensure_ascii=False),
                          financeiro_sync_em=datetime.now().isoformat(timespec="minutes"), financeiro_sync_erro="")
            conn.commit()
            return ""
        finally:
            conn.close()


def sincronizar_se_preciso():
    """Atualiza em segundo plano se a última leitura foi há mais de 1 hora (não trava a página)."""
    from database import conectar_db

    if _trava.locked():
        return
    conn = conectar_db()
    configurado, ultima = arquivo_configurado(conn), status(conn)["em"]
    conn.close()
    if configurado and (ultima is None or datetime.now() - ultima > INTERVALO):
        threading.Thread(target=_sincronizar_quieto, daemon=True).start()


def _sincronizar_quieto():
    try:
        sincronizar()
    except Exception:
        pass


# ---------- juntar com os moradores do app ----------

def montar_mes(conn, dados, mes, morador_id):
    """Junta a aba do mês com os moradores do app (pelo apelido ou pelos nomes alternativos)."""
    from modulos.casa import mapa_moradores

    info = dados.get("meses", {}).get(mes)
    if not info:
        return None
    ids = mapa_moradores(conn)
    fotos = {m["id"]: m["foto_versao"] for m in conn.execute("SELECT id, foto_versao FROM moradores")}
    linhas = []
    for linha in info["moradores"]:
        mid = ids.get(normalizar(linha["nome"]))
        linhas.append({**linha, "morador_id": mid, "foto_versao": fotos.get(mid, 0)})
    minha = next((l for l in linhas if l["morador_id"] == morador_id), None)
    divida = None
    for nome, saldo in dados.get("creditos", {}).items():
        if ids.get(normalizar(nome)) == morador_id:
            divida = saldo
    quadros = info["quadros"]
    mensalidades = next((q for q in quadros if "mensal" in normalizar(q["titulo"])),
                        quadros[0] if quadros else {"titulo": "Gastos", "itens": [], "total": 0})
    return {"nome": info["nome"], "linhas": linhas, "minha": minha, "divida": divida,
            "mensalidades": mensalidades, "detalhes": [q for q in quadros if q is not mensalidades and q["itens"]],
            "total_aluguel": sum(l["total"] for l in linhas)}
