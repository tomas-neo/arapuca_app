"""Contas lidas da planilha do financeiro (layout igual ao Financeiro Arapuca 2K26)."""
import importlib
import io
import sys
from pathlib import Path

import openpyxl
import pytest
from werkzeug.security import generate_password_hash

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from modulos import financeiro  # noqa: E402


def planilha_exemplo():
    wb = openpyxl.Workbook()
    resumo = wb.active
    resumo.title = "Resumo"
    resumo["A1"] = "FINANCEIRO 2K24 - REPÚBLICA ARAPUCA"   # o Resumo ficou com o ano velho
    for nome_aba, extra in (("Setembro", 0), ("Outubro", 1000)):
        ws = wb.create_sheet(nome_aba)
        ws["A1"] = "FINANCEIRO 2K26 - REPÚBLICA ARAPUCA"
        cab = {"B": "Quarto", "C": "Ponderação", "F": "Morador", "G": "Casa", "H": "Despesas", "I": "Outros",
               "J": "Pendências", "K": "Caixinha", "L": "Desconto", "M": "Aluguel", "N": "Pagou", "O": "Saldo",
               "R": "Despesas"}
        for col, txt in cab.items():
            ws[f"{col}5"] = txt
        linhas = [("Suíte", "Zangado", 393.07, 428.25, 871.32 + extra / 100, 871.32),
                  (None, "Chuverin", 393.07, 428.25, 871.32, 0),
                  ("Motel", "Oca", 360.61, 428.25, 838.86, 100),
                  (None, "VAGO", 360.61, 428.25, 838.86, 0)]
        for i, (quarto, nome, casa, desp, total, pagou) in enumerate(linhas, start=6):
            if quarto:
                ws[f"B{i}"] = quarto
            ws[f"F{i}"], ws[f"G{i}"], ws[f"H{i}"], ws[f"K{i}"] = nome, casa, desp, 50
            ws[f"M{i}"], ws[f"N{i}"], ws[f"O{i}"] = total, pagou, pagou - total
        ws["F12"] = "TOTAL"
        ws["R7"], ws["Z7"] = "Mensalidades", "Detalhamento Gastos Cartão de Crédito"
        ws["R8"], ws["U8"], ws["Z8"] = "TOTAL", 7000, "TOTAL"
        for col, txt in zip("RSTU", ["Data", "Descrição", "Status", "Valor"]):
            ws[f"{col}9"] = txt
        for col, txt in zip(["Z", "AA", "AB", "AC", "AD"], ["Data", "Descrição", "Parcela", "Responsável", "Valor"]):
            ws[f"{col}9"] = txt
        for i, (desc, status, valor) in enumerate([("Aluguel", "Pago", 6086.25), ("Luz", "Pendente", 693.79)], start=10):
            ws[f"S{i}"], ws[f"T{i}"], ws[f"U{i}"] = desc, status, valor
        ws["S13"], ws["U13"] = "Internet", 163.24          # depois de uma linha vazia
        ws["T21"], ws["U21"] = "TOTAL despesas já pagas", 0
        ws["Z10"], ws["AA10"], ws["AC10"], ws["AD10"] = "24/09", "Temperos", "Chuveiro", 30.54
        ws["AE10"] = "#REF!"
    cred = wb.create_sheet("Créditos")
    cred["B5"], cred["C5"] = "Morador", "Saldo"
    for i, (nome, saldo) in enumerate([("Zangado", 0), ("Chuverin", -150.5), ("Cabeça Oca", -824.99)], start=6):
        cred[f"B{i}"], cred[f"C{i}"] = nome, saldo
    cred["B9"], cred["C9"] = "TOTAL", -975.49
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def test_le_a_planilha():
    dados = financeiro.ler_arquivo(planilha_exemplo())
    assert dados["ano"] == 2026 and sorted(dados["meses"]) == ["2026-09", "2026-10"]
    out = dados["meses"]["2026-10"]
    zangado = out["moradores"][0]
    assert zangado["nome"] == "Zangado" and zangado["quarto"] == "Suíte"
    assert zangado["despesas"] == 42825 and zangado["total"] == 88132 and zangado["caixinha"] == 5000
    assert out["moradores"][1]["quarto"] == "Suíte" and [m["nome"] for m in out["moradores"]][-1] == "VAGO"
    mensal, cartao = out["quadros"]
    assert mensal["titulo"] == "Mensalidades" and [i["descricao"] for i in mensal["itens"]] == ["Aluguel", "Luz", "Internet"]
    assert mensal["itens"][0]["status"] == "Pago" and mensal["total"] == 694328
    assert cartao["itens"] == [{"descricao": "Temperos", "valor": 3054, "data": "24/09", "status": "",
                                "parcela": "", "responsavel": "Chuveiro"}]
    assert dados["creditos"] == {"Zangado": 0, "Chuverin": -15050, "Cabeça Oca": -82499}
    assert financeiro.ler_arquivo(planilha_exemplo(), ano=2027)["ano"] == 2027


def test_id_do_link():
    assert financeiro.id_do_link("https://docs.google.com/spreadsheets/d/1tcOZL25ky2tpNEXGexLTUy7OzJjWRJVE/edit#gid=0") \
        == "1tcOZL25ky2tpNEXGexLTUy7OzJjWRJVE"
    assert financeiro.id_do_link("https://drive.google.com/open?id=1tcOZL25ky2tpNEXGexLTUy7OzJjWRJVE") \
        == "1tcOZL25ky2tpNEXGexLTUy7OzJjWRJVE"
    assert financeiro.id_do_link("qualquer coisa") == ""


@pytest.fixture
def app(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setenv("ARAPUCA_SECRET_KEY", "teste")
    (tmp_path / "dados_iniciais").mkdir()
    (tmp_path / "dados_iniciais" / "moradores.json").write_text(
        '{"Chuveiro": {"nomes": ["Chuverin"]}, "Oca": {"nomes": ["Cabeça Oca"]}}', encoding="utf-8")
    monkeypatch.setattr(financeiro, "baixar", lambda _id: planilha_exemplo())
    sys.modules.pop("app", None)
    modulo = importlib.import_module("app")
    from database import conectar_db
    conn = conectar_db()
    senha = generate_password_hash("segredo1")
    conn.executemany("INSERT INTO moradores (apelido, senha_hash, admin, cargos) VALUES (?, ?, ?, ?)",
                     [("Oca", senha, 1, ""), ("Zangado", senha, 0, "financeiro"), ("Chuveiro", senha, 0, "")])
    conn.commit()
    conn.close()
    return modulo.app


def logar(app, apelido):
    c = app.test_client()
    c.post("/login", data={"apelido": apelido, "senha": "segredo1"})
    return c


def test_financeiro_liga_a_planilha_e_todos_veem(app, monkeypatch):
    link = "https://docs.google.com/spreadsheets/d/1tcOZL25ky2tpNEXGexLTUy7OzJjWRJVE/edit"
    assert logar(app, "Chuveiro").post("/contas/planilha", data={"link": link}).status_code == 403
    zangado = logar(app, "Zangado")
    r = zangado.post("/contas/planilha", data={"link": link, "ano": "2026"})
    assert "Planilha+ligada" in r.headers["Location"]

    pagina = logar(app, "Chuveiro").get("/contas/?mes=2026-10").text
    assert "R$ 871,32" in pagina and "Sua dívida com a rep" in pagina and "R$ 150,50" in pagina   # Chuverin
    assert "Temperos" in pagina and "Internet" in pagina and "VAGO" in pagina
    assert "Lançar" not in pagina and "Ligar planilha" not in pagina   # ninguém lança no app
    assert "Atualizar agora" in zangado.get("/contas/?mes=2026-10").text

    pagina = logar(app, "Oca").get("/contas/?mes=2026-09").text   # Cabeça Oca = Oca
    assert "R$ 838,86" in pagina and "R$ 824,99" in pagina and "Pagou R$ 100,00" in pagina
    assert "?mes=2026-10" in pagina   # navega só pelos meses que existem na planilha


def test_falha_mantem_os_ultimos_dados(app, monkeypatch):
    zangado = logar(app, "Zangado")
    zangado.post("/contas/planilha", data={"link": "1tcOZL25ky2tpNEXGexLTUy7OzJjWRJVE"})

    def sem_internet(_id):
        raise OSError("sem internet")
    monkeypatch.setattr(financeiro, "baixar", sem_internet)
    r = zangado.post("/contas/planilha/atualizar", data={"mes": "2026-10"})
    assert "sem+internet" in r.headers["Location"]
    pagina = zangado.get("/contas/?mes=2026-10").text
    assert "R$ 881,32" in pagina and "Última leitura falhou" in pagina

    zangado.post("/contas/planilha/desligar")
    assert "Ligar planilha" in zangado.get("/contas/").text   # volta pro modo do app
